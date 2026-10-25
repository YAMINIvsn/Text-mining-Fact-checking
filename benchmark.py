"""
benchmark.py - run the fact checker on many videos and save everything to Excel.

Usage:  python benchmark.py urls.txt results.xlsx [--no-cache]
        (urls.txt = one YouTube URL per line; '#' lines ignored)

Then open results.xlsx, fill the 'gold_label' column (dropdown) in the Claims
sheet, and the Metrics sheet updates itself.
"""
import sys
import time

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from factcheck import pipeline

try:
    from factcheck.transcript import extract_video_id
except ImportError:
    from factcheck.transcript.youtube import extract_video_id

COOLDOWN = 60  # seconds between videos that needed fresh API calls (free tiers)

HEADERS = ["video_id", "claim_id", "claim_text", "category", "time", "system_verdict",
           "confidence", "confidence_without_nli", "nli_agreement", "nli_adjustment",
           "conflict_capped", "n_sources", "summary", "cited_urls", "gold_label", "correct"]


def get_segments(video_id, use_cache):
    cached = pipeline.get_cached_transcript(video_id, use_cache)
    if cached:
        return cached["segments"]
    segments, method = pipeline.fetch_official_transcript(video_id), "official"
    if not segments:
        segments, method = pipeline.fetch_whisper_transcript(video_id), "whisper"
    pipeline.save_transcript(video_id, segments, method)
    return segments


def analyse(video_id, use_cache):
    segments = get_segments(video_id, use_cache)
    claims = pipeline.get_cached_claims(video_id, use_cache)
    if claims is None:
        claims = pipeline.fetch_claims(segments)
        pipeline.save_claims(video_id, claims)
    return claims, pipeline.get_claim_verdicts(claims, use_cache)


def without_nli(bd):
    """Confidence the original formula would give (NLI adjustment removed)."""
    if not bd or "weighted_sum" not in bd:
        return None
    base = bd["weighted_sum"]
    if bd.get("conflict_capped"):
        base = min(base, 0.60)
    return round(max(0.0, min(0.95, base)), 3)


def rows_for(video_id, claims, results):
    rows = []
    for c in claims:
        r = results[c.claim_id]
        v = r["verdict"]
        bd = v.confidence_breakdown or {}
        nli = bd.get("nli") or {}
        base = without_nli(bd)
        rows.append([
            video_id, c.claim_id, c.claim_text, c.category, f"{c.start_time}-{c.end_time}",
            v.verdict, v.confidence, base if base is not None else v.confidence,
            nli.get("agreement"), nli.get("adjustment"), bd.get("conflict_capped"),
            len(r["sources"]), v.summary, "\n".join(s.url for s in v.cited_sources),
            None, None,
        ])
    return rows


def write_book(all_rows, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Claims"
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for i, row in enumerate(all_rows, start=2):
        ws.append(row)
        ws[f"P{i}"] = f'=IF(O{i}="","",IF(O{i}&""=F{i}&"",1,0))'
    dv = DataValidation(type="list", formula1='"TRUE,FALSE,UNVERIFIABLE"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"O2:O{max(len(all_rows) + 1, 200)}")
    for col, w in {"C": 60, "M": 50, "N": 40}.items():
        ws.column_dimensions[col].width = w

    m = wb.create_sheet("Metrics")
    m["A1"], m["B1"] = "Accuracy (labelled claims)", '=IFERROR(SUM(Claims!P:P)/COUNT(Claims!P:P),"")'
    m.append([])
    m.append(["Class", "Predicted", "Gold", "TP", "Precision", "Recall"])
    for i, cls in enumerate(["TRUE", "FALSE", "UNVERIFIABLE"], start=4):
        m[f"A{i}"] = cls
        m[f"B{i}"] = f'=SUMPRODUCT((Claims!$F$2:$F$500&""=A{i})*(Claims!$O$2:$O$500&""<>""))'
        m[f"C{i}"] = f'=SUMPRODUCT(--(Claims!$O$2:$O$500&""=A{i}))'
        m[f"D{i}"] = f'=SUMPRODUCT((Claims!$F$2:$F$500&""=A{i})*(Claims!$O$2:$O$500&""=A{i}))'
        m[f"E{i}"] = f'=IFERROR(D{i}/B{i},"")'
        m[f"F{i}"] = f'=IFERROR(D{i}/C{i},"")'
    m["A8"], m["B8"], m["C8"] = "Mean confidence", "With NLI", "Without NLI"
    for r, (label, flag) in zip((9, 10), (("Correct verdicts", 1), ("Incorrect verdicts", 0))):
        m[f"A{r}"] = label
        m[f"B{r}"] = f'=IFERROR(AVERAGEIFS(Claims!G:G,Claims!P:P,{flag}),"")'
        m[f"C{r}"] = f'=IFERROR(AVERAGEIFS(Claims!H:H,Claims!P:P,{flag}),"")'
    m["A11"], m["B11"], m["C11"] = "Gap (correct - incorrect)", '=IFERROR(B9-B10,"")', '=IFERROR(C9-C10,"")'
    m.column_dimensions["A"].width = 28
    wb.save(path)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    use_cache = "--no-cache" not in sys.argv
    urls_file, out = (args + ["urls.txt", "results.xlsx"][len(args):])[:2]

    with open(urls_file, encoding="utf-8") as f:
        urls = [u.strip() for u in f if u.strip() and not u.startswith("#")]

    all_rows = []
    for n, url in enumerate(urls, 1):
        vid = extract_video_id(url)
        if not vid:
            print(f"[{n}/{len(urls)}] skipped (bad URL): {url}")
            continue
        print(f"[{n}/{len(urls)}] {vid}")
        needs_api = not pipeline.is_fully_cached(vid, use_cache)
        for attempt in range(1, 6):
            try:
                claims, results = analyse(vid, use_cache)
                all_rows += rows_for(vid, claims, results)
                write_book(all_rows, out)  # save after every video
                print(f"    {len(claims)} claims -> {out}")
                break
            except Exception as e:
                if "429" in str(e) or "rate_limit" in str(e):
                    print(f"    rate limit, waiting 65s (attempt {attempt}/5)")
                    time.sleep(65)  # finished claims are cached, so the retry resumes
                else:
                    print(f"    FAILED: {e}")
                    break
        if needs_api and n < len(urls):
            time.sleep(COOLDOWN)
    print(f"Done. Open {out}, fill the gold_label column, then check the Metrics sheet.")


if __name__ == "__main__":
    main()